"""Baut aus einer Skript-Markdown-Datei die Word-Fassung.
Gesprochener Text steht linksbündig, Zeilen mit '>>' (Schnitt/Edit) eingerückt, grau, kursiv.
Aufruf: python3 baue_word.py "skripte/Skripte 1 bis 30.md"
"""
import re, sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor

LABELS = {"VIRAL-VORBILD", "COVER-TITEL", "VISUELLE HOOK", "EMOTION", "SKRIPT", "CAPTION"}
src = sys.argv[1]
doc = Document()
doc.styles["Normal"].font.size = Pt(11)
first = True
for raw in open(src, encoding="utf-8").read().split("\n"):
    line = raw.rstrip()
    if line in ("", "---"):
        continue
    if line.startswith("# "):
        doc.add_heading(line[2:], 0)
    elif line.startswith("### "):
        if not first:
            doc.add_page_break()
        first = False
        doc.add_heading(line[4:], 1)
    elif line.startswith(">>"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(2.5)
        r = p.add_run("✂ " + line[2:].strip())
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    elif line.split(" (")[0] in LABELS:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run(line)
        r.bold = True
    else:
        p = doc.add_paragraph()
        text = re.sub(r"`", "", line)
        if text.startswith("Hinweis:"):
            p.add_run(text).italic = True
        else:
            p.add_run(text)
doc.save(re.sub(r"\.md$", ".docx", src))
